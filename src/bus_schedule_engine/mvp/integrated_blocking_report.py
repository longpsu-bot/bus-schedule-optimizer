"""Three inspectable native Excel tables plus exact JSON witnesses (additive exports)."""

from __future__ import annotations

import json
from hashlib import sha256
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .integrated_blocking_validation import validate_assignment, validate_request


def export_reports(cases, output_dir):
    """Save one consolidated workbook with three tables; never trust cached metrics."""
    records = []
    for request, result in cases:
        request = validate_request(request)
        result = dict(result)
        if result.get("blocks") is not None:
            result["validation"] = validate_assignment(request, result["blocks"])
            if result["upper_bound"] != len(result["blocks"]):
                raise ValueError("report upper bound differs from validated assignment")
        records.append(dict(request=request, result=result))
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    tables = {
        "FeasibilityReport": [["Case", "Field", "Value", "Evidence / scope"]],
        "Timetable": [
            [
                "Case",
                "Trip",
                "Direction",
                "Origin",
                "Destination",
                "Departure min",
                "Arrival min",
                "Runtime min",
                "Vehicle",
                "Team",
                "Driver",
                "Attendant",
                "Policy SHA256",
                "Source pointer",
                "Reason",
            ]
        ],
        "VehicleCrewDuties": [
            [
                "Case",
                "Block",
                "Record",
                "Trip / connection",
                "Metric",
                "Start min",
                "End min",
                "Minutes / value",
                "Terminal",
                "Required",
                "Covered",
                "Detail",
            ]
        ],
    }
    for record in records:
        r, res = record["request"], record["result"]
        case = f"{r['runtime_scenario']['id']}_{r['crew_policy']['duty']['limit']}"
        checked = res.get("validation")
        summary = tables["FeasibilityReport"]
        for key in (
            "status",
            "native_status",
            "model_complete",
            "validator_pass",
            "lower_bound",
            "upper_bound",
            "absolute_gap",
            "relative_gap",
            "termination_reason",
            "wall_seconds",
            "deterministic_seconds",
            "proof_scope",
            "operational_authority",
            "request_sha256",
            "policy_sha256",
            "runtime_sha256",
            "model_sha256",
            "response_sha256",
        ):
            summary.append(
                [
                    case,
                    key,
                    res.get(key),
                    "Complete fixed timetable and declared experimental policy only",
                ]
            )
        for key, value in res.get("relaxation", {}).items():
            if key not in {"matching", "source_side_cut"}:
                summary.append(
                    [case, key, value, "Independent matching/max-flow cut and workload relaxation"]
                )
        summary.append(
            [case, "policy_authority", r["source"]["authority"], r["source"]["provenance"]]
        )
        summary.append([case, "stock_mode", r["stock"]["mode"], r["stock"]["provenance"]])
        summary.append(
            [
                case,
                "whole_free_interval_protection",
                True,
                "Explicitly selected rest; no observed-rest or legal-compliance claim",
            ]
        )
        for uncertainty in r["operational_evidence"]:
            summary.append([case, "unverified_operational_evidence", None, uncertainty])
        for key, value in r["crew_policy"].items():
            summary.append(
                [
                    case,
                    "crew_policy." + key,
                    json.dumps(value, sort_keys=True),
                    "PO declared experimental policy",
                ]
            )
        assigned = {}
        if checked:
            for key in ("required_fleet", "crew_rules"):
                summary.append([case, key, checked[key], "Independently reconstructed witness"])
            for role in ("teams", "drivers", "attendants"):
                summary.append(
                    [
                        case,
                        role,
                        checked["required_fleet"],
                        "One unique driver and attendant per vehicle block",
                    ]
                )
            for terminal, count in checked["initial_vehicles_required"].items():
                summary.append(
                    [
                        case,
                        "initial_required_" + terminal,
                        count,
                        "Exact for selected assignment; not a separately proved terminal minimum",
                    ]
                )
            for b in checked["blocks"]:
                for trip in b["trip_ids"]:
                    assigned[trip] = b
            for meal in ("lunch", "dinner"):
                summary.append(
                    [
                        case,
                        meal + "_required",
                        sum(b["meals"][meal]["required"] for b in checked["blocks"]),
                        "Eligibility independently recomputed",
                    ]
                )
                summary.append(
                    [
                        case,
                        meal + "_required_covered",
                        sum(
                            b["meals"][meal]["required"] and b["meals"][meal]["covered"]
                            for b in checked["blocks"]
                        ),
                        "Selected window-contained meal intervals",
                    ]
                )
            for field in ("required", "covered"):
                summary.append(
                    [
                        case,
                        "general_rest_" + field,
                        sum(
                            b["general_rest"]["required"]
                            and (b["general_rest"]["covered"] if field == "covered" else True)
                            for b in checked["blocks"]
                        ),
                        "Independent of meal-window applicability",
                    ]
                )
        physical = (
            {t["trip_id"]: t for t in checked["physical_assignment"]["trips"]} if checked else {}
        )
        for t in r["timetable"]["trips"]:
            b = assigned.get(t["trip_id"], {})
            arrival = physical[t["trip_id"]]["arrival"] if checked else None
            tables["Timetable"].append(
                [
                    case,
                    t["trip_id"],
                    t["direction"],
                    "A" if t["direction"] == "outbound" else "B",
                    "B" if t["direction"] == "outbound" else "A",
                    t["departure"],
                    arrival,
                    arrival - t["departure"] if arrival is not None else None,
                    *[b.get(k) for k in ("vehicle_id", "team_id", "driver_id", "attendant_id")],
                    res.get("policy_sha256"),
                    t["source_pointer"],
                    None if checked else "no validated incumbent",
                ]
            )
        detail = tables["VehicleCrewDuties"]
        for b in checked["blocks"] if checked else []:
            prefix = [case, b["block_id"]]
            for metric in (
                "duty_span",
                "total_driving",
                "max_continuous_driving",
                "continuous_work_exposure",
                "protected_minutes",
            ):
                detail.append(
                    prefix
                    + [
                        "BLOCK",
                        None,
                        metric,
                        b["duty_start"] if metric == "duty_span" else None,
                        b["duty_end"] if metric == "duty_span" else None,
                        b[metric],
                        b["initial_terminal"],
                        None,
                        None,
                        "; ".join(
                            str(b[k])
                            for k in (
                                "vehicle_id",
                                "team_id",
                                "driver_id",
                                "attendant_id",
                                "final_terminal",
                            )
                        ),
                    ]
                )
            for index, trip in enumerate(b["trip_ids"], 1):
                detail.append(
                    prefix + ["TRIP", trip, "trip_order", None, None, index, None, None, None, None]
                )
            for m, data in b["meals"].items():
                detail.append(
                    prefix
                    + [
                        "COVERAGE",
                        None,
                        m,
                        None,
                        None,
                        data["overlap"],
                        None,
                        data["required"],
                        data["covered"],
                        "meal-window overlap minutes; meal subintervals are rest labels",
                    ]
                )
                for start, end in data["intervals"]:
                    detail.append(
                        prefix
                        + [
                            "MEAL",
                            None,
                            m,
                            start,
                            end,
                            end - start,
                            None,
                            None,
                            None,
                            "Contained in one protected connection interval",
                        ]
                    )
            g = b["general_rest"]
            detail.append(
                prefix
                + [
                    "COVERAGE",
                    None,
                    "general_rest",
                    None,
                    None,
                    len(g["qualifying_intervals"]),
                    None,
                    g["required"],
                    g["covered"],
                    "Requires one uninterrupted qualifying protected interval",
                ]
            )
            for start, end in b["rest_union"]:
                detail.append(
                    prefix
                    + [
                        "REST",
                        None,
                        "protected_union",
                        start,
                        end,
                        end - start,
                        None,
                        None,
                        None,
                        "Physical elapsed rest counted once, including meal/reset sharing",
                    ]
                )
            for start, end in b["reset_intervals"]:
                detail.append(
                    prefix
                    + [
                        "RESET",
                        None,
                        "driving_reset",
                        start,
                        end,
                        end - start,
                        None,
                        None,
                        None,
                        "Label on same physical rest; not additional elapsed time",
                    ]
                )
            for c in b["connections"]:
                edge = c["after_trip"] + " -> " + c["before_trip"]
                for metric in (
                    "dwell",
                    "technical",
                    "post",
                    "pre",
                    "reserve",
                    "protected_minutes",
                    "other_wait",
                ):
                    detail.append(
                        prefix
                        + [
                            "CONNECTION",
                            edge,
                            metric,
                            c["free_start"],
                            c["free_end"],
                            c[metric],
                            c["terminal"],
                            None,
                            None,
                            "Dwell = technical + post + pre + reserve + protected union + other wait",
                        ]
                    )
            for a in b["activities"]:
                detail.append(
                    prefix
                    + [
                        "ACTIVITY",
                        a["trip_id"],
                        a["kind"],
                        a["start"],
                        a["end"],
                        a["end"] - a["start"],
                        None,
                        None,
                        None,
                        "Disjoint duty activity; technical and reserve only internally",
                    ]
                )
    for name, rows in tables.items():
        sheet = workbook.create_sheet(name)
        if len(rows) == 1:
            rows.append(["No validated duty rows"] + [None] * (len(rows[0]) - 1))
        for row in rows:
            sheet.append(
                [json.dumps(v, sort_keys=True) if isinstance(v, (list, dict)) else v for v in row]
            )
        sheet.freeze_panes = "C2"
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[1].height = 32
        for cell in sheet[1]:
            cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="173E59")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for cells in sheet.iter_rows(min_row=2):
            height = 18
            for cell in cells:
                cell.font = Font(name="Calibri", size=11, color="243B4A")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, str):
                    height = max(height, min(180, 15 * ((len(cell.value) // 65) + 1)))
                if isinstance(cell.value, int):
                    cell.number_format = "0"
                elif isinstance(cell.value, float):
                    cell.number_format = "0.000"
            sheet.row_dimensions[cells[0].row].height = height
        for i, head in enumerate(rows[0], 1):
            width = 18
            if head in {"Field", "Metric", "Trip / connection", "Trip"}:
                width = 30
            elif head in {"Evidence / scope", "Detail", "Source pointer"}:
                width = 90
            elif head.endswith("SHA256"):
                width = 24
            elif head == "Value":
                width = 55
            sheet.column_dimensions[get_column_letter(i)].width = width
        table = Table(
            displayName=name, ref=f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
        )
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
        sheet.print_title_rows = "1:1"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
    xlsx = output_dir / "integrated_vehicle_crew_v1.xlsx"
    workbook.save(xlsx)
    destination = output_dir / "integrated_vehicle_crew_v1.json"
    payload = dict(
        schema_version="integrated_blocking_report_v1",
        cases=records,
        analysis_code_sha256={
            p.name: sha256(p.read_bytes().replace(b"\r\n", b"\n")).hexdigest()
            for p in Path(__file__).parent.glob("*integrated_blocking*.py")
        },
    )
    destination.write_text(
        json.dumps(payload, sort_keys=True, indent=2, allow_nan=False) + "\n", encoding="utf-8"
    )
    return dict(xlsx=str(xlsx.resolve()), json=str(destination.resolve()))
