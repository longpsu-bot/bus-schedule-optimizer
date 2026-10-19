"""Bounded feedback coordinator and operator-reviewable JSON/XLSX outputs."""

from __future__ import annotations

import json
from collections import Counter
from copy import deepcopy
from itertools import product
from pathlib import Path
from time import perf_counter as monotonic

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .circulation import assign_fleet
from .input import DIRECTIONS, integer, number, validate_bundle
from .joint import local_evidence, pair_key, pair_priority
from .planning import (
    _compiled_settlement,
    compile_plan,
    initial_plan,
    revise_plan,
    revision_priority,
)
from .quality import change_justification, evaluate, regime_failures, select_candidate
from .review import build_review, warning_policy
from .review_html import render_review_html

DEFAULTS = {
    "candidate_budget": 25,
    "max_revisions": 24,
    "solver_seconds": 2,
    "solver_work_limit": 0.2,
    "total_seconds": 120,
    "minimum_internal_gaps": 2,
    "minimum_regime_minutes": 30,
    "maximum_regimes": 4,
    "boundary_settlement_minutes": 60,
    "compilation_lookahead": 0,
    "directional_settlements": 1,
    "pair_evaluation_budget": 160,
    "pairs_per_plan": 16,
    "wait_absolute_tolerance": 0.10,
    "wait_relative_tolerance": 0.02,
    "wait_materiality_minutes": 0.05,
    "wait_materiality_relative": 0.01,
}


def _settings(settings):
    if settings is not None and not isinstance(settings, dict):
        raise ValueError("settings must be an object")
    active = DEFAULTS | (settings or {})
    if set(active) - (set(DEFAULTS) | {"maximum_transition_jump"}):
        raise ValueError("unknown planning setting")
    for key in ("candidate_budget", "minimum_internal_gaps", "maximum_regimes"):
        integer(active[key], key, 1)
    if active["minimum_internal_gaps"] < 2:
        raise ValueError("minimum_internal_gaps must be >=2; one-gap regimes are forbidden")
    for key in ("max_revisions", "minimum_regime_minutes", "boundary_settlement_minutes"):
        integer(active[key], key)
    integer(active["compilation_lookahead"], "compilation_lookahead")
    if active["compilation_lookahead"] > 256:
        raise ValueError("compilation_lookahead must be <=256")
    for key in ("directional_settlements", "pair_evaluation_budget", "pairs_per_plan"):
        integer(active[key], key, 1)
    if active["directional_settlements"] > 4:
        raise ValueError("directional_settlements must be <=4")
    if active["pairs_per_plan"] > 16:
        raise ValueError("pairs_per_plan must be <=16")
    for key in ("solver_seconds", "solver_work_limit", "total_seconds"):
        number(active[key], key)
        if active[key] <= 0:
            raise ValueError(f"{key} must be positive")
    for key in (
        "wait_absolute_tolerance",
        "wait_relative_tolerance",
        "wait_materiality_minutes",
        "wait_materiality_relative",
    ):
        number(active[key], key)
    if active.get("maximum_transition_jump") is not None:
        integer(active["maximum_transition_jump"], "maximum_transition_jump")
    return active


def run_bundle(
    bundle, settings=None, *, review_settings=None, operating_scenarios=None, incumbents=None
):
    result = {
        "status": "BLOCKED_INPUT",
        "search_status": "BLOCKED_INPUT",
        "operator_decision": "UNRESOLVED",
        "selected_source": None,
        "recommended": None,
        "alternatives": [],
        "diagnostics": {"evaluations": 0, "feedback": []},
    }
    try:
        data, active = validate_bundle(bundle), _settings(settings)
        warning_policy(review_settings)
        if incumbents is not None:
            if (
                not isinstance(incumbents, list)
                or len(incumbents) > 1
                or active["directional_settlements"] == 1
            ):
                raise ValueError(
                    "incumbents requires joint mode and at most one supplied timetable"
                )
            for c in incumbents:
                if not isinstance(c, dict) or not {"departures", "regimes"} <= set(c):
                    raise ValueError("incumbents requires departures and regimes")
                if not isinstance(c["departures"], dict) or set(c["departures"]) != set(DIRECTIONS):
                    raise ValueError("incumbents requires both directional departure vectors")
                for d in DIRECTIONS:
                    times = c["departures"].get(d, [])
                    direction = data["directions"][d]
                    if (
                        not isinstance(times, list)
                        or len(times) != direction["trip_count"]
                        or any(type(t) is not int for t in times)
                        or times != sorted(set(times))
                        or times[0] != direction["first"]
                        or times[-1] != direction["last"]
                    ):
                        raise ValueError("incumbents must preserve exact whole-minute authority")
                if not isinstance(c["regimes"], list) or not c["regimes"]:
                    raise ValueError("incumbents requires uniform ServiceRegime evidence")
                if regime_failures(c["departures"], c["regimes"], active):
                    raise ValueError("incumbents contains invalid ServiceRegime evidence")
                if any(
                    not isinstance(r.get("regime_type"), str)
                    or r["regime_type"] not in {"normal", "peak", "low-demand", "end_tail"}
                    for r in c["regimes"]
                ):
                    raise ValueError("incumbents requires supported regime_type metadata")
    except ValueError as exc:
        result["diagnostics"]["input_error"] = str(exc)
        return result
    result.update(
        route_id=data["route_id"],
        provenance=data["provenance"],
        settings=active,
        demand_basis=data["demand_basis"],
    )
    for direction in DIRECTIONS:
        d = data["directions"][direction]
        if d["trip_count"] > d["last"] - d["first"] + 1:
            result["status"] = "PROVEN_INFEASIBLE"
            result["search_status"] = "PROVEN_INFEASIBLE"
            result["diagnostics"]["proof"] = (
                f"{direction}: {d['trip_count']} trips exceed {d['last'] - d['first'] + 1} distinct whole-minute slots between authoritative endpoints"
            )
            return result
    deadline = monotonic() + active["total_seconds"]
    initial = {d: initial_plan(data["directions"][d], active) for d in DIRECTIONS}
    initial_key = json.dumps(initial, sort_keys=True)
    pending = {
        initial_key: {
            "operation": "initial",
            "plans": initial,
            "direction": None,
            "priority": (4, 0),
            "guided": False,
            "sequence": 0,
        }
    }
    evaluated, generated = set(), {initial_key}
    direction_counts, operation_counts, rejections = Counter(), Counter(), Counter()
    generated_count, max_pending, compiled_count, duplicates, guided_count = 1, 1, 0, 0, 0
    schedules = set()
    retained_pair_keys = set()
    joint = active["directional_settlements"] > 1
    compilation_cache, pair_cache = {}, {}
    exact_pairs = compilation_hits = 0
    pair_limited = False
    unexamined_pairs = 0
    candidates = []
    limited = False
    baseline = None
    baseline_times = {d: data["directions"][d]["baseline"] for d in DIRECTIONS}
    result["baseline"] = {}
    if monotonic() < deadline:
        baseline_fleet = assign_fleet(data, baseline_times, settings=active, deadline=deadline)
        if monotonic() < deadline and baseline_fleet["status"] != "SOLVER_LIMIT":
            result["baseline"] = evaluate(data, baseline_times, [], baseline_fleet, active)
            baseline = {
                "id": "B",
                "source": "BASELINE",
                "departures": baseline_times,
                "regimes": [],
                "vehicle_chains": baseline_fleet,
                "quality": result["baseline"],
            }
            result["baseline_evidence"] = baseline
    baseline_incomplete = (
        not baseline or baseline["quality"]["operational_certification"] == "INCOMPLETE"
    )
    comparison_incomplete = baseline_incomplete and not result["baseline"].get(
        "proven_hard_failures"
    )
    result["diagnostics"]["baseline_status"] = (
        "TIME_LIMIT" if not baseline else "INCOMPLETE" if baseline_incomplete else "EVALUATED"
    )
    max_evaluations = min(25, active["candidate_budget"], active["max_revisions"] + 1)
    incumbent_evidence = []
    for supplied in incumbents or []:
        if monotonic() >= deadline:
            limited = True
            break
        times, regimes = deepcopy(supplied["departures"]), deepcopy(supplied["regimes"])
        # Canonicalize equal-headway labels before a complete candidate is cached by vector.
        canonical_regimes = []
        for d in DIRECTIONS:
            blocks = [r for r in regimes if r["direction"] == d]
            settled = _compiled_settlement(
                data["directions"][d],
                [{"end": r["end_time"], "gaps": r["gap_count"]} for r in blocks],
                [r["headway_minutes"] for r in blocks],
            )
            canonical_regimes.extend(r | {"direction": d} for r in settled["regimes"])
        regimes = canonical_regimes
        bounded = active | {"solver_seconds": min(active["solver_seconds"], deadline - monotonic())}
        fleet = assign_fleet(data, times, settings=bounded, deadline=deadline)
        candidate = {
            "id": "I001",
            "source": "OPTIMIZED",
            "verified_incumbent": True,
            "departures": times,
            "regimes": regimes,
            "vehicle_chains": fleet,
            "quality": evaluate(data, times, regimes, fleet, bounded),
        }
        local_evidence(data, candidate, baseline, review_settings, active)
        key = pair_key(data, times, active, review_settings)
        schedules.add(key)
        retained_pair_keys.add(key)
        pair_cache[key] = candidate
        exact_pairs += 1
        incumbent_evidence.append(candidate)
        if candidate["quality"]["hard_constraints_pass"]:
            candidates.append(candidate)
    solver_limited = baseline_incomplete
    quality_feedback = []
    while pending:
        if (
            pair_limited
            or result["diagnostics"]["evaluations"] >= max_evaluations
            or monotonic() >= deadline
        ):
            limited = True
            break
        if initial_key in pending and not evaluated:
            key = initial_key
        else:
            available = {entry["direction"] for entry in pending.values()}
            direction = min(available, key=lambda d: (direction_counts[d], DIRECTIONS.index(d)))
            pool = [k for k, entry in pending.items() if entry["direction"] == direction]
            # Every third turn explores a less-used operation; other turns exploit measured feedback.
            diverse = direction_counts[direction] % 3 == 2
            key = min(
                pool,
                key=lambda k: (
                    operation_counts[(direction, pending[k]["operation"].split(":")[1])]
                    if diverse
                    else 0,
                    pending[k]["priority"],
                    pending[k]["sequence"],
                    k,
                ),
            )
        entry = pending.pop(key)
        evaluated.add(key)
        operation, plans = entry["operation"], entry["plans"]
        if entry["direction"] is not None:
            direction_counts[entry["direction"]] += 1
            operation_counts[(entry["direction"], operation.split(":")[1])] += 1
        guided_count += entry["guided"]
        result["diagnostics"]["evaluations"] += 1
        index = result["diagnostics"]["evaluations"]
        feedback = {
            "evaluation": index,
            "operation": operation,
            "plans": plans,
            "compilation": {},
            "demand_guided": entry["guided"],
            "revision_priority": list(entry["priority"]),
        }
        departures, regimes, directional = {}, [], {}
        for direction in DIRECTIONS:
            remaining = deadline - monotonic()
            if remaining <= 0:
                feedback["compilation"][direction] = {
                    "status": "SOLVER_LIMIT",
                    "reason": "total run seconds reached",
                }
                limited = True
                break
            bounded = active | {"solver_seconds": min(active["solver_seconds"], remaining)}
            compilation_key = json.dumps([direction, plans[direction]], sort_keys=True)
            cached_compile = joint and compilation_key in compilation_cache
            if cached_compile:
                compiled = deepcopy(compilation_cache[compilation_key])
                compilation_hits += 1
            else:
                compiled = compile_plan(
                    data["directions"][direction],
                    plans[direction],
                    bounded,
                    **({"deadline": deadline} if joint else {}),
                )
                if joint:
                    compilation_cache[compilation_key] = deepcopy(compiled)
            feedback["compilation"][direction] = {
                k: v
                for k, v in compiled.items()
                if k not in {"departures", "regimes", "settlements"}
            }
            if joint and compiled.get("settlement_lookahead_complete") is False:
                solver_limited = True
            if joint:
                feedback["compilation"][direction]["cache_hit"] = cached_compile
                if cached_compile:
                    feedback["compilation"][direction]["settlement_previews"] = 0
            if compiled["status"] != "COMPILED":
                rejections[compiled["status"]] += 1
                solver_limited |= compiled["status"] == "SOLVER_LIMIT"
                continue
            departures[direction] = compiled["departures"]
            directional[direction] = compiled.get("settlements", [compiled])
            average = (
                data["directions"][direction]["last"] - data["directions"][direction]["first"]
            ) / (data["directions"][direction]["trip_count"] - 1)
            for settlement in directional[direction]:
                for r in settlement["regimes"]:
                    r["direction"] = direction
                    if r["regime_type"] != "end_tail":
                        r["regime_type"] = (
                            "peak" if r["headway_minutes"] < average else "low-demand"
                        )
        reason = "COMPILATION"
        quality = None
        if len(departures) == 2:
            compiled_count += 1
            pairs = list(product(*(directional[d] for d in DIRECTIONS)))
            if joint:
                retained_pair_keys.update(
                    pair_key(
                        data,
                        {d: c["departures"] for d, c in zip(DIRECTIONS, p, strict=True)},
                        active,
                        review_settings,
                    )
                    for p in pairs
                )
            if joint:
                pairs.sort(key=lambda p: pair_priority(p, quality_feedback))
            unexamined_pairs += max(0, len(pairs) - active["pairs_per_plan"])
            pair_records = []
            for pair_index, pair in enumerate(pairs[: active["pairs_per_plan"]], 1):
                departures = {d: c["departures"] for d, c in zip(DIRECTIONS, pair, strict=True)}
                regimes = [r for c in pair for r in c["regimes"]]
                key = (
                    pair_key(data, departures, active, review_settings)
                    if joint
                    else json.dumps(departures, sort_keys=True)
                )
                record = {"pair_index": pair_index}
                if key in schedules:
                    record.update(candidate_status="DUPLICATE_TIMETABLE")
                    if joint:
                        record["cached_candidate_id"] = pair_cache[key]["id"]
                        record["quality"] = pair_cache[key]["quality"]
                    duplicates += 1
                    pair_records.append(record)
                    continue
                if exact_pairs >= active["pair_evaluation_budget"]:
                    pair_limited = limited = True
                    unexamined_pairs += len(pairs[: active["pairs_per_plan"]]) - pair_index + 1
                    break
                remaining = deadline - monotonic()
                if remaining <= 0:
                    limited = True
                    unexamined_pairs += len(pairs[: active["pairs_per_plan"]]) - pair_index + 1
                    break
                schedules.add(key)
                exact_pairs += 1
                bounded = active | {"solver_seconds": min(active["solver_seconds"], remaining)}
                fleet = assign_fleet(data, departures, settings=bounded, deadline=deadline)
                quality = evaluate(data, departures, regimes, fleet, bounded)
                candidate = {
                    "id": f"C{index:03d}_{pair_index:02d}" if joint else f"C{index:03d}",
                    "source": "OPTIMIZED",
                    "departures": departures,
                    "regimes": regimes,
                    "vehicle_chains": fleet,
                    "quality": quality,
                }
                if joint:
                    local_evidence(data, candidate, baseline, review_settings, active)
                    pair_cache[key] = candidate
                record.update(
                    candidate_status="FEASIBLE" if quality["hard_constraints_pass"] else "REJECTED",
                    fleet_status=fleet["status"],
                    failures=quality["failures"],
                    quality=quality,
                    departures=departures,
                    regimes=regimes,
                )
                if joint:
                    record["candidate"] = candidate
                pair_records.append(record)
                solver_limited |= fleet["status"] in {"SOLVER_LIMIT", "MINIMUM_FLEET_UNPROVED"}
                rejections.update(quality["failures"])
                if quality["hard_constraints_pass"]:
                    candidates.append(candidate)
            if pair_records:
                measured = [p for p in pair_records if p.get("quality")]
                representative = (
                    min(
                        measured,
                        key=lambda p: (
                            not p["quality"]["hard_constraints_pass"],
                            p["quality"]["required_fleet"]
                            if p["quality"]["required_fleet"] is not None
                            else float("inf"),
                            p["quality"]["waiting_minutes"] or 0,
                            p["pair_index"],
                        ),
                    )
                    if measured
                    else pair_records[0]
                )
                feedback.update({k: v for k, v in representative.items() if k != "candidate"})
                quality = representative.get("quality")
                if quality:
                    reason = (
                        "QUALITY"
                        if quality["hard_constraints_pass"]
                        else "FLEET"
                        if any("FLEET" in x for x in quality["failures"])
                        else "SERVICE"
                    )
                    quality_feedback = quality["service_feedback"]
            if joint:
                feedback["pairs"] = pair_records
                feedback["retained_pair_count"] = len(pairs)
        result["diagnostics"]["feedback"].append(feedback)
        for direction in DIRECTIONS:
            evidence = (
                [f for f in quality["service_feedback"] if f["direction"] == direction]
                if quality
                else reason
            )
            for name, changed in revise_plan(
                data["directions"][direction], plans[direction], evidence, active
            ):
                generated_count += 1
                revised = plans | {direction: changed}
                key = json.dumps(revised, sort_keys=True)
                generated.add(key)
                if key in evaluated:
                    continue
                priority = revision_priority(
                    data["directions"][direction], plans[direction], name, changed, evidence
                )
                if quality is None and direction not in departures:
                    priority = (4, -len(departures))
                proposal = {
                    "operation": f"{direction}:{name}",
                    "plans": revised,
                    "direction": direction,
                    "priority": priority,
                    "guided": priority[0] < 4,
                    "sequence": generated_count,
                }
                if key not in pending or priority < pending[key]["priority"]:
                    pending[key] = proposal
            # Keep an independently bounded pool for each direction; evicted plans may re-enter later.
            keys = sorted(
                [k for k, p in pending.items() if p["direction"] == direction],
                key=lambda k: (pending[k]["priority"], pending[k]["sequence"], k),
            )
            for k in keys[2 * max_evaluations :]:
                del pending[k]
        max_pending = max(max_pending, len(pending))
    if joint:
        unexamined_pairs = len(retained_pair_keys - schedules)
    limited = limited or solver_limited or unexamined_pairs > 0
    result["recommended"] = select_candidate(candidates, active, baseline) if baseline else None
    result["search_status"] = (
        ("FEASIBLE_BUDGET_LIMITED" if limited else "FEASIBLE_RECOMMENDATION")
        if candidates
        else "SEARCH_INCOMPLETE_NO_FEASIBLE_CANDIDATE"
    )
    selected = result["recommended"]
    result["selected_source"] = selected["source"] if selected else None
    result["operator_decision"] = (
        "KEEP_BASELINE"
        if selected and selected["source"] == "BASELINE"
        else "REPLACE_BASELINE"
        if selected
        else "UNRESOLVED"
    )
    result["status"] = (
        "NO_CHANGE_RECOMMENDED"
        if result["operator_decision"] == "KEEP_BASELINE"
        else "BASELINE_COMPARISON_INCOMPLETE"
        if candidates and comparison_incomplete
        else result["search_status"]
    )
    result["decision"] = (
        change_justification(selected, baseline, active)
        if selected and selected["source"] == "OPTIMIZED" and baseline
        else {
            "justified": False,
            "reason": "NO_JUSTIFIED_REPLACEMENT_WITHIN_SEARCH"
            if selected
            else "NO_CERTIFIED_TIMETABLE",
            "baseline_hard_constraints_pass": result["baseline"].get("hard_constraints_pass"),
        }
    )
    if candidates:
        alternatives = [c for c in candidates if not selected or c["id"] != selected["id"]]
        while alternatives and len(result["alternatives"]) < 2:
            selected = select_candidate(alternatives, active)
            selected["change_assessment"] = (
                change_justification(selected, baseline, active)
                if baseline
                else {"justified": False, "reason": "BASELINE_CERTIFICATION_UNAVAILABLE"}
            )
            result["alternatives"].append(selected)
            alternatives.remove(selected)
        sse_anchor = min(candidates, key=lambda c: (c["quality"]["sse"], c["id"]))["id"]
        te_anchor = min(candidates, key=lambda c: (c["quality"]["te"], c["id"]))["id"]
        result["diagnostics"]["metric_disagreement"] = {
            "sse_best": sse_anchor,
            "te_best": te_anchor,
            "disagree": sse_anchor != te_anchor,
            "selection_gate": False,
        }
    if not result["recommended"]:
        result["diagnostics"]["action"] = (
            "Inspect structured family/compiler/stock/protection feedback; increase bounded budget or revise service-plan family. Route infeasibility is not proved."
        )
    rejected = [
        f for f in result["diagnostics"]["feedback"] if f.get("candidate_status") == "REJECTED"
    ]
    best_rejected = min(
        rejected,
        key=lambda f: (
            any("FLEET" in failure for failure in f["failures"]),
            any("PROTECTED" in failure for failure in f["failures"]),
            max(
                (
                    x["severity"]
                    for x in f["quality"]["service_feedback"]
                    if x["problem"] == "PROTECTED_SERVICE"
                ),
                default=0,
            ),
            max(
                (
                    x["severity"]
                    for x in f["quality"]["service_feedback"]
                    if x["problem"] == "CAPACITY_PROXY"
                ),
                default=0,
            ),
            f["evaluation"],
        ),
        default=None,
    )
    result["diagnostics"].update(
        budget_limited=limited,
        feasible_candidates=len(candidates),
        feasible_generated_candidates=sum(not c.get("verified_incumbent") for c in candidates),
        incumbents=incumbent_evidence,
        distinct_plans_generated=len(generated),
        total_plans_generated=generated_count,
        distinct_plans_evaluated=len(evaluated),
        compiled_candidates=compiled_count,
        duplicate_schedules=duplicates,
        exact_pair_evaluations=exact_pairs,
        pair_evaluation_budget=active["pair_evaluation_budget"],
        directional_compilation_cache_hits=compilation_hits,
        unexamined_retained_pairs=unexamined_pairs,
        distinct_retained_pairs_discovered=len(retained_pair_keys) if joint else len(schedules),
        rejections_by_reason=dict(sorted(rejections.items())),
        demand_guided_evaluations=guided_count,
        fallback_evaluations=result["diagnostics"]["evaluations"] - guided_count,
        directional_evaluations=dict(direction_counts),
        maximum_pending_plans=max_pending,
        pending_plans=len(pending),
        search_exhausted=not pending and not solver_limited and not unexamined_pairs,
        search_limit_reason="TOTAL_DEADLINE"
        if monotonic() >= deadline
        else "PAIR_EVALUATION_BUDGET"
        if pair_limited
        else "EVALUATION_BUDGET"
        if limited and pending
        else "SOLVER_LIMIT"
        if solver_limited
        else "PAIRS_PER_PLAN"
        if unexamined_pairs
        else None,
        remaining_service_violations=[
            f
            for f in (
                result["recommended"]["quality"]
                if result["recommended"]
                else best_rejected["quality"]
                if best_rejected
                else result["baseline"]
            ).get("service_feedback", [])
            if f["hard"]
        ],
        best_rejected_evaluation=best_rejected["evaluation"] if best_rejected else None,
        best_rejected_failures=best_rejected["failures"] if best_rejected else [],
        best_rejected_service_violations=[
            f for f in best_rejected["quality"]["service_feedback"] if f["hard"]
        ]
        if best_rejected
        else [],
        selection_policy=(
            "hard admissibility > explicit benefit versus B > verified-incumbent non-regression > protected service > wait-equivalence > maximum material local deterioration > positive passenger-minute deterioration > review warning severity/count and gaps > fleet > regimes > transitions > terminal wait > tail"
            if joint
            else "hard admissibility > explicit benefit versus B > protected service > wait-equivalence > fleet > regimes > transitions > terminal wait > tail; secondary metrics never justify change alone"
        ),
        global_optimality_proved=False,
    )
    if joint:
        selected = result["recommended"]
        result["diagnostics"]["joint_selection"] = {
            "selected_id": selected["id"] if selected else None,
            "selected_local_service": selected.get("local_service") if selected else None,
            "verified_incumbent_retained": bool(selected and selected.get("verified_incumbent")),
            "minimum_feasible_fleet": min(
                (c["quality"]["required_fleet"] for c in candidates), default=None
            ),
            "best_feasible_waiting_minutes": min(
                (c["quality"]["waiting_minutes"] or 0 for c in candidates), default=None
            ),
            "baseline_exact_evaluations": int(baseline is not None),
            "pair_budget_includes_incumbent_revalidation": True,
            "unexamined_pair_basis": "Distinct retained vectors discovered so far, excluding every evaluated or cached vector; uncompiled plan families are uncounted.",
        }
    result["operational_review"] = build_review(data, result, review_settings, operating_scenarios)
    return result


def export_result(result, directory):
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    recommended = result["recommended"]
    report = {
        "status": result["status"],
        "search_status": result.get("search_status", result["status"]),
        "operator_decision": result.get("operator_decision", "UNRESOLVED"),
        "selected_source": result.get("selected_source"),
        "decision": result.get("decision"),
        "route_id": result.get("route_id"),
        "quality": recommended["quality"] if recommended else None,
        "baseline": result.get("baseline"),
        "provenance": result.get("provenance"),
        "settings": result.get("settings"),
        "demand_basis": result.get("demand_basis"),
        "operational_review": result.get("operational_review"),
        "alternatives": [
            {
                "id": c["id"],
                "source": c["source"],
                "quality": c["quality"],
                "change_assessment": c.get("change_assessment"),
            }
            for c in result["alternatives"]
        ],
    }
    outputs = {
        "recommended_timetable.json": {
            "status": result["status"],
            "departures": recommended["departures"] if recommended else None,
            "candidate_id": recommended["id"] if recommended else None,
            "selected_source": result.get("selected_source"),
            "search_status": result.get("search_status", result["status"]),
            "operator_decision": result.get("operator_decision", "UNRESOLVED"),
        },
        "service_regimes.json": recommended["regimes"] if recommended else [],
        "vehicle_chains.json": recommended["vehicle_chains"] if recommended else None,
        "quality_report.json": report,
        "run_diagnostics.json": result["diagnostics"],
    }
    for name, payload in outputs.items():
        (directory / name).write_text(
            json.dumps(payload, ensure_ascii=False, sort_keys=True, indent=2, allow_nan=False)
            + "\n",
            encoding="utf-8",
        )
    workbook = Workbook()
    workbook.remove(workbook.active)
    for direction in DIRECTIONS:
        sheet = workbook.create_sheet(direction.title())
        sheet.append(["Trip", "Departure minute", "Departure HH:MM", "Arrival HH:MM", "Vehicle"])
        if recommended:
            rows = sorted(
                (t for t in recommended["vehicle_chains"]["trips"] if t["direction"] == direction),
                key=lambda t: t["departure"],
            )
            for t in rows:
                sheet.append(
                    [
                        t["trip_id"],
                        t["departure"],
                        _time(t["departure"]),
                        _time(t["arrival"]),
                        t["vehicle_id"],
                    ]
                )
    sheet = workbook.create_sheet("ServiceRegimes")
    sheet.append(
        [
            "Direction",
            "Start",
            "End",
            "Headway minutes",
            "Planned departures",
            "Internal gaps",
            "Type",
        ]
    )
    for r in recommended["regimes"] if recommended else []:
        sheet.append(
            [
                r["direction"],
                _time(r["start_time"]),
                _time(r["end_time"]),
                r["headway_minutes"],
                r["planned_departure_count"],
                r["gap_count"],
                r["regime_type"],
            ]
        )
    sheet = workbook.create_sheet("VehicleCirculation")
    sheet.append(["Vehicle", "Trip", "Direction", "Departure", "Arrival", "Ready after layover"])
    for t in recommended["vehicle_chains"]["trips"] if recommended else []:
        sheet.append(
            [
                t["vehicle_id"],
                t["trip_id"],
                t["direction"],
                _time(t["departure"]),
                _time(t["arrival"]),
                _time(t["ready"]),
            ]
        )
    sheet = workbook.create_sheet("QualityComparison")
    sheet.append(["Measure", "Existing B", "Recommendation"])
    sheet.append(["Run status", "independently evaluated", result["status"]])
    sheet.append(["Selected source", "BASELINE", result.get("selected_source")])
    sheet.append(["Search status", None, result.get("search_status")])
    sheet.append(["Operator decision", None, result.get("operator_decision")])
    sheet.append(["Decision evidence", None, json.dumps(result.get("decision"))])
    for metric in (
        "hard_constraints_pass",
        "waiting_minutes",
        "waiting_difference_minutes",
        "waiting_difference_percent",
        "required_fleet",
        "maximum_protected_gap",
        "regime_count",
        "transition_jump",
        "terminal_wait",
        "tail_gap",
        "maximum_headway",
        "maximum_transition_jump",
        "protected_service_certified",
        "protection_basis",
        "sse",
        "te",
        "continuous_exposure",
        "capacity",
        "directional_waiting_minutes",
        "outside_service_passengers",
    ):
        baseline = result.get("baseline", {}).get(metric)
        value = recommended["quality"].get(metric) if recommended else None
        sheet.append(
            [
                metric,
                json.dumps(baseline) if isinstance(baseline, dict) else baseline,
                json.dumps(value) if isinstance(value, dict) else value,
            ]
        )
    sheet.append(
        [
            "Waiting assumption",
            "Uniform bucket arrivals",
            "Modeled scheduled wait, not observed wait",
        ]
    )
    sheet = workbook.create_sheet("DemandBuckets")
    columns = [
        "direction",
        "interval_start",
        "interval_end",
        "demand_passengers",
        "baseline_departures",
        "candidate_departures",
        "waiting_minutes",
        "baseline_waiting_minutes",
        "waiting_difference_minutes",
        "waiting_difference_percent",
        "waiting_integral",
        "waiting_contribution_minutes",
        "maximum_intersecting_gap",
        "maximum_modeled_boardings",
        "capacity_planning_threshold",
        "protected_service_violations",
        "nearby_regimes",
    ]
    sheet.append(["Source", *columns])
    for source, quality in [
        ("BASELINE", result.get("baseline", {})),
        *[
            (c["id"], c["quality"])
            for c in ([recommended] if recommended and recommended["source"] == "OPTIMIZED" else [])
            + result["alternatives"]
        ],
    ]:
        for row in quality.get("demand_buckets", []):
            sheet.append(
                [
                    source,
                    *[
                        json.dumps(row.get(k))
                        if isinstance(row.get(k), (list, dict))
                        else row.get(k)
                        for k in columns
                    ],
                ]
            )
    sheet = workbook.create_sheet("Alternatives")
    sheet.append(
        [
            "Candidate",
            "Source",
            "Hard constraints pass",
            "Fleet",
            "Modeled wait",
            "Difference from B",
            "Justified",
            "Reason",
        ]
    )
    for candidate in result["alternatives"]:
        q, decision = candidate["quality"], candidate.get("change_assessment", {})
        sheet.append(
            [
                candidate["id"],
                candidate["source"],
                q["hard_constraints_pass"],
                q["required_fleet"],
                q["waiting_minutes"],
                q["waiting_difference_minutes"],
                decision.get("justified"),
                decision.get("reason"),
            ]
        )
    sheet = workbook.create_sheet("ServiceFeedback")
    sheet.append(
        [
            "Evaluation",
            "Operation",
            "Direction",
            "Start minute",
            "End minute",
            "Problem",
            "Hard",
            "Severity",
            "Units",
            "Evidence",
        ]
    )
    for evaluation in result["diagnostics"].get("feedback", []):
        for f in evaluation.get("quality", {}).get("service_feedback", []):
            sheet.append(
                [
                    evaluation["evaluation"],
                    evaluation["operation"],
                    f["direction"],
                    f["interval_start"],
                    f["interval_end"],
                    f["problem"],
                    f["hard"],
                    f["severity"],
                    f["severity_units"],
                    json.dumps(f),
                ]
            )
    sheet.append(
        [
            "Demand basis",
            result.get("demand_basis"),
            "Historical boardings are an arrival proxy"
            if result.get("demand_basis") == "historical_boardings"
            else "Independent arrivals",
        ]
    )
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="183B56")
            cell.font = Font(color="FFFFFF", bold=True)
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(
                70, max(14, max(len(str(c.value or "")) for c in column) + 2)
            )
    workbook.save(directory / "recommended_timetable.xlsx")
    (directory / "operational_review.html").write_text(
        render_review_html(result.get("operational_review")), encoding="utf-8"
    )


def _time(minutes):
    return f"{minutes // 60:02d}:{minutes % 60:02d}"
