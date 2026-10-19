"""Deterministic adapter for the verified eight-sheet external report layout."""

from __future__ import annotations

import hashlib
import math
import re
from datetime import datetime, time, timedelta
from pathlib import Path

import openpyxl

TABLES = {
    "BASELINE": {"outbound": "hien_huu_From_MienTay", "inbound": "hien_huu_From_DHQG-A."},
    "EXTERNAL_CANDIDATE": {
        "outbound": "xanh_sm_From_MienTay",
        "inbound": "xanh_sm_From_DHQG-A.",
    },
}
SHEETS = [
    "Tong Quan",
    *[name for pair in TABLES.values() for name in pair.values()],
    "Lich cho tai xe",
    "So chuyen cho tung xe",
    "Du lieu bieu do",
]


def _number(value, location, *, integer=False):
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"{location}: expected a nonnegative number")
    if not math.isfinite(value) or value < 0 or (integer and value != int(value)):
        raise ValueError(f"{location}: malformed nonnegative number")
    return int(value) if integer else value


def _time(value, location, *, whole=True):
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        minute = value.hour * 60 + value.minute + value.second / 60 + value.microsecond / 6e7
    elif isinstance(value, timedelta):
        minute = value.total_seconds() / 60
    elif isinstance(value, (float, int)) and not isinstance(value, bool) and 0 <= value < 1:
        minute = value * 1440
    elif isinstance(value, str) and re.fullmatch(r"\d{2}:\d{2}(?::\d{2})?", value):
        h, m, *s = map(int, value.split(":"))
        if h > 23 or m > 59 or (s and s[0] > 59):
            raise ValueError(f"{location}: malformed timestamp")
        minute = h * 60 + m + (s[0] / 60 if s else 0)
    else:
        raise ValueError(f"{location}: malformed timestamp")
    if not 0 <= minute < 1440 or (whole and not math.isclose(minute, round(minute), abs_tol=1e-7)):
        raise ValueError(f"{location}: timestamp must be a whole minute in the operating day")
    return int(round(minute)) if whole else minute


def _timetable(sheet, read, summary):
    hours, total_col, total_row, rows = [], None, None, []
    for col in range(2, sheet.max_column + 1):
        cell = sheet.cell(1, col)
        value = read(sheet.title, cell.coordinate)
        if value == "Tổng":
            if total_col is not None:
                raise ValueError(f"{sheet.title}: duplicate total column")
            total_col = col
        elif isinstance(value, str) and re.fullmatch(r"(?:[01]\d|2[0-3])h", value):
            hours.append((col, int(value[:2])))
        elif value is not None or any(
            sheet.cell(r, col).value is not None for r in range(2, sheet.max_row + 1)
        ):
            raise ValueError(f"{sheet.title}!{cell.coordinate}: malformed hour column")
    if not hours or [h for _, h in hours] != sorted({h for _, h in hours}) or total_col is None:
        raise ValueError(f"{sheet.title}: missing, duplicate or unsorted hour columns")
    for row in range(2, sheet.max_row + 1):
        value = read(sheet.title, f"A{row}")
        if value == "Tổng cộng" and total_row is None:
            total_row = row
        elif (
            isinstance(value, str) and re.fullmatch(r"Chuyến [1-9]\d*", value) and total_row is None
        ):
            if int(value.split()[-1]) != len(rows) + 1:
                raise ValueError(f"{sheet.title}!A{row}: conflicting trip row sequence")
            rows.append(row)
        elif any(sheet.cell(row, c).value is not None for c in range(1, sheet.max_column + 1)):
            raise ValueError(f"{sheet.title}!A{row}: unsupported or conflicting timetable row")
    if total_row is None:
        raise ValueError(f"{sheet.title}: missing Tổng cộng row")
    trips, hourly_totals = [], []
    for col, hour in hours:
        local, blank_seen = [], False
        for row in rows:
            cell = sheet.cell(row, col)
            value = read(sheet.title, cell.coordinate)
            if value is None:
                blank_seen = True
                continue
            if blank_seen:
                raise ValueError(f"{sheet.title}!{cell.coordinate}: missing preceding minute entry")
            if isinstance(value, str) and re.fullmatch(r"\d{1,2}", value):
                value = int(value)
            minute = _number(value, f"{sheet.title}!{cell.coordinate}", integer=True)
            if minute > 59:
                raise ValueError(f"{sheet.title}!{cell.coordinate}: minute outside 0..59")
            local.append(hour * 60 + minute)
            trips.append(
                {
                    "departure": local[-1],
                    "source_sheet": sheet.title,
                    "source_cells": [sheet.cell(1, col).coordinate, cell.coordinate],
                }
            )
        if local != sorted(set(local)):
            raise ValueError(f"{sheet.title}: duplicate or unsorted departures")
        count_cell = sheet.cell(total_row, col).coordinate
        count = _number(read(sheet.title, count_cell), f"{sheet.title}!{count_cell}", integer=True)
        if count != len(local):
            raise ValueError(
                f"{sheet.title}!{count_cell}: hourly total differs from individual entries"
            )
        hourly_totals.append({"hour": hour, "source_cell": count_cell, "count": count})
    count_cell = sheet.cell(total_row, total_col).coordinate
    count = _number(read(sheet.title, count_cell), f"{sheet.title}!{count_cell}", integer=True)
    if not trips or len(trips) != count or count != summary["trip_count"]:
        raise ValueError(f"{sheet.title}: directional total conflicts with entries or summary")
    if (
        trips[0]["departure"] != summary["first_departure"]
        or trips[-1]["departure"] != summary["last_departure"]
    ):
        raise ValueError(f"{sheet.title}: first/last departure conflicts with summary")
    return trips, {
        "source_sheet": sheet.title,
        "hourly_totals": hourly_totals,
        "directional_total": {"source_cell": count_cell, "count": count},
    }


def _demand(sheet, read):
    buckets, exclusions = [], []
    for coordinate, expected in {
        "A1": "Giờ",
        "D1": "Khách/ngày",
        "F1": "Giờ",
        "I1": "Khách/ngày",
        "K1": "Giờ",
        "L1": "Trong khung",
        "M1": "Ngoài khung",
        "O1": "Giờ",
        "P1": "Trong khung",
        "Q1": "Ngoài khung",
    }.items():
        if read(sheet.title, coordinate) != expected:
            raise ValueError(f"{sheet.title}!{coordinate}: unsupported demand header")
    for direction, tc, dc, rt, rv in [("outbound", 1, 4, 11, 12), ("inbound", 6, 9, 15, 16)]:
        raw = []
        for row in range(2, sheet.max_row + 1):
            tcell, vcell = sheet.cell(row, rt), sheet.cell(row, rv)
            t, v = read(sheet.title, tcell.coordinate), read(sheet.title, vcell.coordinate)
            excluded_cell = sheet.cell(row, rv + 1).coordinate
            excluded = read(sheet.title, excluded_cell)
            if t is None and v is None and excluded is None:
                continue
            raw.append(
                (
                    _time(t, f"{sheet.title}!{tcell.coordinate}"),
                    _number(v, f"{sheet.title}!{vcell.coordinate}"),
                    [tcell.coordinate, vcell.coordinate],
                )
            )
            if excluded is not None:
                excluded = _number(excluded, f"{sheet.title}!{excluded_cell}")
                if excluded:
                    exclusions.append(
                        {
                            "direction": direction,
                            "observation_time": raw[-1][0],
                            "demand_passengers": excluded,
                            "source_sheet": sheet.title,
                            "source_cells": [tcell.coordinate, excluded_cell],
                            "normalization_method": "SOURCE_OUTSIDE_FRAME_EXCLUDED_FROM_HALF_HOUR_SERIES",
                        }
                    )
        if not raw or [t for t, _, _ in raw] != sorted({t for t, _, _ in raw}):
            raise ValueError(f"{direction}: raw in-frame observations missing or not unique/sorted")
        local = []
        display_rows = [
            r
            for r in range(2, sheet.max_row + 1)
            if any(
                read(sheet.title, sheet.cell(r, c).coordinate) is not None
                for c in range(tc, dc + 1)
            )
        ]
        if (
            not display_rows
            or display_rows != list(range(2, display_rows[-1] + 1))
            or len(display_rows) % 2
        ):
            raise ValueError(f"{direction}: incomplete repeated display rows")
        for row in display_rows[::2]:
            tcell, dcell = sheet.cell(row, tc), sheet.cell(row, dc)
            start = _time(read(sheet.title, tcell.coordinate), f"{sheet.title}!{tcell.coordinate}")
            demand = _number(
                read(sheet.title, dcell.coordinate), f"{sheet.title}!{dcell.coordinate}"
            )
            cells = [
                tcell.coordinate,
                sheet.cell(row + 1, tc).coordinate,
                dcell.coordinate,
                sheet.cell(row + 1, dc).coordinate,
            ]
            if (
                read(sheet.title, sheet.cell(row + 1, tc).coordinate) is not None
                or read(sheet.title, cells[-1]) != demand
            ):
                raise ValueError(
                    f"{direction}: repeated chart display pair does not agree at {dcell.coordinate}"
                )
            # The separate in-frame observations establish the bucket convention, not equality alone.
            observations = [(v, refs) for t, v, refs in raw if start <= t < start + 30]
            if not observations or not math.isclose(
                sum(v for v, _ in observations), demand, abs_tol=1e-8
            ):
                raise ValueError(
                    f"{direction}: half-hour total does not reconcile at {dcell.coordinate}"
                )
            if local and start != local[-1]["interval_end"]:
                raise ValueError(f"{direction}: half-hour intervals are not contiguous")
            local.append(
                {
                    "direction": direction,
                    "interval_start": start,
                    "interval_end": start + 30,
                    "demand_passengers": demand,
                    "source_sheet": sheet.title,
                    "source_cells": cells + [c for _, refs in observations for c in refs],
                    "normalization_method": "REPEATED_HALF_HOUR_DISPLAY_RECONCILED_TO_IN_FRAME_OBSERVATIONS",
                }
            )
        if any(not local[0]["interval_start"] <= t < local[-1]["interval_end"] for t, _, _ in raw):
            raise ValueError(f"{direction}: observations outside chart interval coverage")
        buckets.extend(local)
    return buckets, exclusions


def _vehicles(sheet, summaries, read, issues):
    trips, reported = [], []
    headers = [
        "trip_no",
        "vehicle_id",
        "direction",
        "start_time",
        "arrival_time",
        "Thoi gian nghi tai ben",
        "Tai trong du kien",
    ]
    if any(
        read(sheet.title, sheet.cell(1, c + offset).coordinate) != header
        for offset in (0, 8)
        for c, header in enumerate(headers, 1)
    ):
        issues.append(
            {
                "code": "UNSUPPORTED_VEHICLE_HEADERS",
                "source_sheet": sheet.title,
                "source_cells": ["A1:G1", "I1:O1"],
            }
        )
        return trips, reported
    for direction, offset, text in [
        ("outbound", 0, "MienTay --> DHQG-A."),
        ("inbound", 8, "DHQG-A. --> MienTay"),
    ]:
        for row in range(2, sheet.max_row + 1):
            cells = [sheet.cell(row, offset + col).coordinate for col in range(1, 8)]
            values = [read(sheet.title, cell) for cell in cells]
            if all(v is None for v in values):
                continue
            try:
                sequence, vehicle, label, departure, arrival, layover, exposure = values
                if not isinstance(vehicle, str) or not vehicle.strip() or label != text:
                    raise ValueError("missing vehicle or conflicting direction")
                trip = {
                    "vehicle_id": vehicle,
                    "direction": direction,
                    "trip_sequence": _number(sequence, cells[0], integer=True),
                    "departure": _time(departure, cells[3]),
                    "arrival": _time(arrival, cells[4]),
                    "recorded_layover": _number(layover, cells[5]) if layover is not None else None,
                    "reported_boarding_exposure": _number(exposure, cells[6])
                    if exposure is not None
                    else None,
                    "source_sheet": sheet.title,
                    "source_cells": cells,
                }
                if trip["arrival"] <= trip["departure"] or trip["trip_sequence"] < 1:
                    raise ValueError("arrival must follow departure in the same operating day")
                trips.append(trip)
            except ValueError as error:
                issues.append(
                    {
                        "code": "INCOMPLETE_VEHICLE_ROW",
                        "source_sheet": sheet.title,
                        "source_cells": cells,
                        "detail": str(error),
                    }
                )
    if [read(summaries.title, summaries.cell(1, c).coordinate) for c in range(1, 9)] != [
        "vehicle",
        "vehicle_id",
        "n_trips",
        "n_inbound",
        "n_outbound",
        "first_departure",
        "last_arrival",
        "service_min",
    ]:
        issues.append(
            {
                "code": "UNSUPPORTED_VEHICLE_SUMMARY_HEADERS",
                "source_sheet": summaries.title,
                "source_cells": ["A1:H1"],
            }
        )
        return trips, reported
    for row in range(2, summaries.max_row + 1):
        cells = [summaries.cell(row, col).coordinate for col in range(1, 9)]
        values = [read(summaries.title, c) for c in cells]
        if all(v is None for v in values):
            continue
        try:
            _, vehicle, count, inbound, outbound, first, last, duration = values
            if not isinstance(vehicle, str) or not vehicle.strip():
                raise ValueError("missing vehicle identifier")
            reported.append(
                {
                    "vehicle_id": vehicle,
                    "trip_count": _number(count, cells[2], integer=True),
                    "inbound_count": _number(inbound, cells[3], integer=True),
                    "outbound_count": _number(outbound, cells[4], integer=True),
                    "first_departure": _time(first, cells[5], whole=False),
                    "last_arrival": _time(last, cells[6], whole=False),
                    "reported_operating_duration": _number(duration, cells[7]),
                    "source_sheet": summaries.title,
                    "source_cells": cells,
                }
            )
        except ValueError as error:
            issues.append(
                {
                    "code": "INCOMPLETE_VEHICLE_SUMMARY",
                    "source_sheet": summaries.title,
                    "source_cells": cells,
                    "detail": str(error),
                }
            )
    return trips, reported


def import_external_workbook(path):
    path = Path(path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    workbook = openpyxl.load_workbook(path, data_only=False)
    cached = openpyxl.load_workbook(path, data_only=True)
    formula_cells, issues = set(), []

    def read(sheet, coordinate):
        cell = workbook[sheet][coordinate]
        value = cell.value
        if cell.data_type == "f":
            value = cached[sheet][coordinate].value
            formula_cells.add(f"{sheet}!{coordinate}")
            if value is None:
                raise ValueError(
                    f"{sheet}!{coordinate}: formula has no cached value; recalculate in source application"
                )
        if cell.data_type == "e":
            raise ValueError(f"{sheet}!{coordinate}: spreadsheet error {value}")
        return value

    try:
        if set(workbook.sheetnames) != set(SHEETS):
            raise ValueError("expected the verified eight-sheet external report layout")
        summary, timetables, controls = {}, {}, {}
        for source, pair in TABLES.items():
            summary[source], timetables[source], controls[source] = {}, {}, {}
            for direction, sheet in pair.items():
                col = {
                    ("BASELINE", "outbound"): "C",
                    ("EXTERNAL_CANDIDATE", "outbound"): "D",
                    ("BASELINE", "inbound"): "J",
                    ("EXTERNAL_CANDIDATE", "inbound"): "K",
                }[source, direction]
                item = {
                    "trip_count": _number(
                        read("Tong Quan", f"{col}3"), f"Tong Quan!{col}3", integer=True
                    ),
                    "first_departure": _time(read("Tong Quan", f"{col}4"), f"Tong Quan!{col}4"),
                    "last_departure": _time(read("Tong Quan", f"{col}5"), f"Tong Quan!{col}5"),
                    "reported_fleet": _number(
                        read("Tong Quan", f"{col}8"), f"Tong Quan!{col}8", integer=True
                    ),
                    "claimed_peak_headway": read("Tong Quan", f"{col}6"),
                    "claimed_offpeak_headway": read("Tong Quan", f"{col}7"),
                    "source_sheet": "Tong Quan",
                    "source_cells": [f"{col}{r}" for r in range(3, 9)],
                }
                summary[source][direction] = item
                timetables[source][direction], controls[source][direction] = _timetable(
                    workbook[sheet], read, item
                )
        try:
            demand, exclusions = _demand(workbook["Du lieu bieu do"], read)
            for direction in ("outbound", "inbound"):
                local = [b for b in demand if b["direction"] == direction]
                if any(
                    local[0]["interval_start"] > timetables[s][direction][0]["departure"]
                    or local[-1]["interval_end"] < timetables[s][direction][-1]["departure"]
                    for s in TABLES
                ):
                    raise ValueError(
                        f"{direction}: demand coverage does not cover the full B/X service horizon"
                    )
        except ValueError as error:
            demand, exclusions = [], []
            issues.append(
                {
                    "code": "UNRESOLVED_DEMAND_AGGREGATION",
                    "source_sheet": "Du lieu bieu do",
                    "detail": str(error),
                }
            )
        vehicle_trips, vehicle_summaries = _vehicles(
            workbook["Lich cho tai xe"], workbook["So chuyen cho tung xe"], read, issues
        )
        route = {"outbound": read("Tong Quan", "A1"), "inbound": read("Tong Quan", "H1")}
        if not all(isinstance(v, str) and v.strip() for v in route.values()):
            raise ValueError("Tong Quan!A1/H1: missing directional route identity")
        return {
            "schema_version": "external_timetable_v1",
            "source": {
                "workbook": path.name,
                "sha256": digest,
                "sheets": workbook.sheetnames,
                "producer_status": "UNVERIFIED",
                "external_label": "Xanh SM",
                "formula_cache_cells": sorted(formula_cells),
            },
            "route_identity": route,
            "operating_day_profile": None,
            "summary": summary,
            "timetables": timetables,
            "timetable_controls": controls,
            "demand": demand,
            "demand_exclusions": exclusions,
            "demand_status": "SUPPORTED_BOARDING_PROXY"
            if demand
            else "UNRESOLVED_DEMAND_AGGREGATION",
            "demand_basis": "In-frame historical boarding/chart observations; not independent passenger-arrival observations. Outside-frame observations are excluded by source convention. Uniform arrivals within each supported half-hour are a modeling assumption.",
            "vehicle_trips": vehicle_trips,
            "vehicle_summaries": vehicle_summaries,
            "issues": issues,
        }
    finally:
        workbook.close()
        cached.close()
