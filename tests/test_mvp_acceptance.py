"""Independent opt-in private-route acceptance; no private inputs in the repository."""

import json
import os
from pathlib import Path

import pytest
from openpyxl import load_workbook

from bus_schedule_engine.mvp.circulation import validate_chains
from bus_schedule_engine.mvp.input import bundle_from_workbook
from bus_schedule_engine.mvp.runner import export_result, run_bundle

ROOT = Path(__file__).resolve().parents[1]


def test_committed_fixture_is_executable(tmp_path):
    data = json.loads((ROOT / "tests/fixtures/mvp/synthetic_route.json").read_text())
    result = run_bundle(data)
    assert result["recommended"]["quality"]["waiting_minutes"] == 10
    assert result["recommended"]["quality"]["required_fleet"] == 2
    export_result(result, tmp_path)
    assert (tmp_path / "recommended_timetable.xlsx").stat().st_size > 1000


def test_workbook_hard_capacity_is_not_dropped_when_source_capacity_is_blank(tmp_path):
    from test_v3_runner import _write_v3_workbook

    path = _write_v3_workbook(tmp_path / "source.xlsx")
    wb = load_workbook(path)
    for row in wb["THONG_SO_B"].iter_rows():
        if row[0].value == "vehicle_capacity_passengers":
            row[1].value = None
    wb.save(path)
    with pytest.raises(ValueError, match="vehicle_capacity"):
        bundle_from_workbook(path, capacity_policy="hard")
    data = bundle_from_workbook(path, capacity_policy="evaluation_only")
    assert data.get("vehicle_capacity") is None


@pytest.mark.parametrize("route", ["6", "10"])
def test_private_route_acceptance_independently(route, tmp_path):
    workbook_path = os.environ.get(f"BUS_MVP_ROUTE{route}_WORKBOOK")
    if not workbook_path:
        pytest.skip(f"BUS_MVP_ROUTE{route}_WORKBOOK is not configured; private data not fabricated")
    data = bundle_from_workbook(workbook_path, minimum_layover=5)
    result = run_bundle(data)
    assert result["route_id"] == route
    candidate = result["recommended"]
    assert candidate is not None, result["diagnostics"]
    assert candidate["quality"]["hard_constraints_pass"]
    assert result["diagnostics"]["evaluations"] <= 25
    assert result == run_bundle(data)
    for direction, times in candidate["departures"].items():
        assert len(times) == data["directions"][direction]["trip_count"]
        assert times[0] == data["directions"][direction]["first"]
        assert times[-1] == data["directions"][direction]["last"]
        assert all(type(t) is int for t in times)
        regimes = [r for r in candidate["regimes"] if r["direction"] == direction]
        if candidate["source"] == "OPTIMIZED":
            assert sum(r["planned_departure_count"] for r in regimes) == len(times)
        else:
            assert times == data["directions"][direction]["baseline"]
            assert regimes == []
            assert result["status"] == "NO_CHANGE_RECOMMENDED"
        for regime in regimes:
            members = [t for t in times if regime["start_time"] <= t <= regime["end_time"]]
            assert all(
                b - a == regime["headway_minutes"]
                for a, b in zip(members, members[1:], strict=False)
            )
            assert regime["gap_count"] >= 2
        assert candidate["quality"]["directional_waiting_minutes"][direction] > 0
    validate_chains(data, candidate["departures"], candidate["vehicle_chains"])
    assert candidate["quality"]["required_fleet"] <= data["fleet_ceiling"]
    export_result(result, tmp_path)
    wb = load_workbook(tmp_path / "recommended_timetable.xlsx", data_only=True)
    for direction in ("outbound", "inbound"):
        assert [
            r[1] for r in wb[direction.title()].iter_rows(min_row=2, values_only=True)
        ] == candidate["departures"][direction]
    assert result["diagnostics"]["metric_disagreement"]["selection_gate"] is False


@pytest.mark.parametrize("route", ["6", "10"])
def test_private_route_strict_proxy_is_independently_certified_or_incomplete(route, tmp_path):
    workbook_path = os.environ.get(f"BUS_MVP_ROUTE{route}_WORKBOOK")
    if not workbook_path:
        pytest.skip(f"BUS_MVP_ROUTE{route}_WORKBOOK is not configured")
    data = bundle_from_workbook(
        workbook_path, minimum_layover=5, maximum_load_factor=0.8, capacity_policy="hard"
    )
    result = run_bundle(data)
    assert result == run_bundle(data)
    assert result["diagnostics"]["evaluations"] <= 25
    assert result["status"] != "PROVEN_INFEASIBLE"
    if result["recommended"]:
        c = result["recommended"]
        assert c["quality"]["hard_constraints_pass"]
        assert c["quality"]["capacity"]["within_policy"]
        assert not c["quality"]["capacity"]["actual_load_factor_certified"]
        validate_chains(data, c["departures"], c["vehicle_chains"])
    else:
        assert result["status"] == "SEARCH_INCOMPLETE_NO_FEASIBLE_CANDIDATE"
        assert not result["baseline"]["hard_constraints_pass"]
        assert result["diagnostics"]["best_rejected_service_violations"]
    export_result(result, tmp_path)
    assert len(list(tmp_path.iterdir())) == 7
