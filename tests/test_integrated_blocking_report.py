import json
import os
import runpy
from pathlib import Path

import pytest
from openpyxl import load_workbook
from test_integrated_blocking_validation import request

from bus_schedule_engine.mvp.integrated_blocking import solve_integrated
from bus_schedule_engine.mvp.integrated_blocking_report import export_reports
from bus_schedule_engine.mvp.route10_integrated_blocking import frozen_rows, route10_requests


def test_three_native_tables_roundtrip_exact_departures_and_linked_duty_rows(tmp_path):
    r = request(
        [("outbound", 480, 660), ("inbound", 705, 885)], meals=True, driving=True, general=True
    )
    result = solve_integrated(r)
    paths = export_reports([(r, result)], tmp_path)
    saved = json.loads(Path(paths["json"]).read_text(encoding="utf-8"))
    assert saved["cases"][0]["result"]["upper_bound"] == 1
    workbook = load_workbook(paths["xlsx"], data_only=False)
    assert workbook.sheetnames == ["FeasibilityReport", "Timetable", "VehicleCrewDuties"]
    assert all(len(s.tables) == 1 for s in workbook)
    timetable = list(workbook["Timetable"].values)
    headers = timetable[0]
    assert [row[headers.index("Departure min")] for row in timetable[1:]] == [480, 705]
    assert [row[headers.index("Arrival min")] for row in timetable[1:]] == [660, 885]
    assert all(row[headers.index("Team")] for row in timetable[1:])
    details = list(workbook["VehicleCrewDuties"].values)
    assert any("general_rest" in row for row in details)
    assert any("protected_union" in row for row in details)


def test_incomplete_report_does_not_invent_assignments_or_arrivals(tmp_path):
    r = request([("outbound", 480, 500)])
    result = solve_integrated(r, cancelled=lambda: True)
    paths = export_reports([(r, result)], tmp_path)
    rows = list(load_workbook(paths["xlsx"], data_only=True)["Timetable"].values)
    head, row = rows
    assert row[head.index("Vehicle")] is None
    assert row[head.index("Arrival min")] is None
    assert row[head.index("Reason")] == "no validated incumbent"


def test_report_revalidates_instead_of_trusting_solver_claims(tmp_path):
    r = request([("outbound", 480, 500)])
    result = solve_integrated(r)
    result["blocks"][0]["trip_ids"] = []
    with pytest.raises(ValueError):
        export_reports([(r, result)], tmp_path)


def test_streaming_frozen_reader_skips_unneeded_fields(tmp_path):
    p = tmp_path / "evidence.json"
    p.write_text(
        json.dumps(
            dict(
                schema_version="route10_fleet_pilot_v1",
                source={"sha256": "x"},
                scenarios=[
                    dict(
                        id="report_p50",
                        bundle={"example": 50},
                        comparison={"C": {"departures": {"outbound": [285]}}},
                        guided={"unneeded": [1] * 1000},
                    ),
                    dict(id="report_p75", bundle={"example": 75}, comparison={"C": {}}, guided={}),
                ],
            ),
            indent=2,
            sort_keys=True,
        ),
        encoding="utf-8",
    )
    rows = frozen_rows(p)
    assert set(rows) == {"report_p50", "report_p75"}
    assert rows["report_p50"]["bundle"] == {"example": 50}
    assert "guided" not in rows["report_p50"]


def test_wrong_private_hash_blocked_before_source_parsing(tmp_path):
    workbook = tmp_path / "private.xlsx"
    workbook.write_bytes(b"wrong source")
    with pytest.raises(ValueError, match="SHA-256"):
        route10_requests(workbook, tmp_path / "missing.json")


def test_cli_protects_source_even_when_named_like_atomic_checkpoint_staging(tmp_path):
    source = tmp_path / "p50_480.checkpoint.json.tmp"
    source.write_bytes(b"must remain unchanged")
    main = runpy.run_path(
        str(Path(__file__).parents[1] / "scripts" / "run_route10_integrated_blocking.py")
    )["main"]
    with pytest.raises(ValueError, match="output/source collision"):
        main(
            [
                "--input",
                str(source),
                "--evidence",
                str(tmp_path / "missing.json"),
                "--output-dir",
                str(tmp_path),
            ]
        )
    assert source.read_bytes() == b"must remain unchanged"


def test_cli_protects_existing_hardlink_to_frozen_source(tmp_path):
    source = tmp_path / "private.xlsx"
    source.write_bytes(b"frozen bytes")
    alias = tmp_path / "p50_480.checkpoint.json.tmp"
    alias.hardlink_to(source)
    main = runpy.run_path(
        str(Path(__file__).parents[1] / "scripts" / "run_route10_integrated_blocking.py")
    )["main"]
    with pytest.raises(ValueError, match="output/source collision"):
        main(
            [
                "--input",
                str(source),
                "--evidence",
                str(tmp_path / "missing.json"),
                "--output-dir",
                str(tmp_path),
            ]
        )
    assert source.read_bytes() == b"frozen bytes"


def test_cli_guards_lexical_staging_sibling_before_resolving_checkpoint_symlink(
    tmp_path, monkeypatch
):
    source = tmp_path / "p50_480.checkpoint.json.tmp"
    source.write_bytes(b"frozen bytes")
    destination = tmp_path / "resume.json"
    destination.write_text("{}", encoding="utf-8")
    alias = tmp_path / "p50_480.checkpoint.json"
    try:
        alias.symlink_to(destination)
    except OSError:
        # Windows may lack symlink privilege. Simulate only the path-resolution
        # boundary; real files and the CLI's collision checks still run.
        original_resolve = Path.resolve

        def resolve(path, *args, **kwargs):
            if path == alias:
                return original_resolve(destination, *args, **kwargs)
            return original_resolve(path, *args, **kwargs)

        monkeypatch.setattr(Path, "resolve", resolve)
    main = runpy.run_path(
        str(Path(__file__).parents[1] / "scripts" / "run_route10_integrated_blocking.py")
    )["main"]
    with pytest.raises(ValueError, match="output/source collision"):
        main(
            [
                "--input",
                str(source),
                "--evidence",
                str(tmp_path / "missing.json"),
                "--output-dir",
                str(tmp_path),
            ]
        )
    assert source.read_bytes() == b"frozen bytes"


@pytest.mark.skipif(
    not os.getenv("BUS_CREW_FROZEN_EVIDENCE"), reason="private frozen input not configured"
)
def test_private_request_adapter_locks_all_102_p0_departures_and_four_bounds():
    requests = route10_requests(
        os.environ["BUS_MVP5_ROUTE10_WORKBOOK"], os.environ["BUS_CREW_FROZEN_EVIDENCE"]
    )
    assert [(r["runtime_scenario"]["id"], r["crew_policy"]["duty"]["limit"]) for r in requests] == [
        ("p50", 480),
        ("p50", 600),
        ("p75", 480),
        ("p75", 600),
    ]
    assert all(len(r["timetable"]["trips"]) == 102 and r["resource_cap"] is None for r in requests)
    results = [solve_integrated(r, cancelled=lambda: True) for r in requests]
    assert [r["relaxation"]["workload_duty_bound"] for r in results] == [26, 21, 27, 22]
    assert [r["relaxation"]["edge_count"] for r in results] == [1987, 1987, 1978, 1978]
    assert [r["relaxation"]["task_path_bound"] for r in results] == [16] * 4
