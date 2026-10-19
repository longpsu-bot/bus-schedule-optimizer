"""Public synthetic adapter contract; no private report is needed in CI."""

from pathlib import Path

import openpyxl
import pytest

from bus_schedule_engine.mvp import external_import

FIXTURE = Path(__file__).parent / "fixtures" / "external_report.xlsx"


def changed(tmp_path, sheet, cell, value):
    workbook = openpyxl.load_workbook(FIXTURE)
    workbook[sheet][cell] = value
    path = tmp_path / "changed.xlsx"
    workbook.save(path)
    workbook.close()
    return path


def test_reconstructs_all_four_51_trip_tables_and_cell_evidence():
    data = external_import.import_external_workbook(FIXTURE)
    for source in ("BASELINE", "EXTERNAL_CANDIDATE"):
        for direction in ("outbound", "inbound"):
            trips = data["timetables"][source][direction]
            assert len(trips) == 51
            assert trips[0]["departure"] == 300
            assert trips[-1]["departure"] == 1050
            assert all(isinstance(t["departure"], int) for t in trips)
            assert trips[0]["source_cells"] == ["C1", "C2"]
    assert data["source"]["producer_status"] == "UNVERIFIED"
    assert len(data["source"]["sha256"]) == 64
    assert data == external_import.import_external_workbook(FIXTURE)


@pytest.mark.parametrize("value", ["60", "-1", 2.5, True, "03:15", "5.0", "=10+5"])
def test_rejects_malformed_minutes_and_uncached_formulas(tmp_path, value):
    path = changed(tmp_path, "hien_huu_From_MienTay", "C2", value)
    with pytest.raises(ValueError, match="C2"):
        external_import.import_external_workbook(path)


@pytest.mark.parametrize(
    ("sheet", "cell", "value"),
    [
        ("hien_huu_From_MienTay", "C3", "00"),
        ("hien_huu_From_MienTay", "C3", "45"),
        ("hien_huu_From_MienTay", "C6", 3),
        ("hien_huu_From_MienTay", "U6", 50),
        ("Tong Quan", "C4", "05:01"),
        ("Tong Quan", "D3", 52),
        ("xanh_sm_From_DHQG-A.", "D1", "05h"),
        ("hien_huu_From_MienTay", "C3", None),
    ],
)
def test_rejects_duplicates_order_conflicts_and_totals(tmp_path, sheet, cell, value):
    with pytest.raises(ValueError):
        external_import.import_external_workbook(changed(tmp_path, sheet, cell, value))


def test_demand_duplicates_reconcile_to_raw_observations_and_zero_is_retained():
    data = external_import.import_external_workbook(FIXTURE)
    assert data["demand_status"] == "SUPPORTED_BOARDING_PROXY"
    buckets = [b for b in data["demand"] if b["direction"] == "outbound"]
    assert len(buckets) == 26
    assert sum(b["demand_passengers"] for b in buckets) == 500
    assert buckets[0]["demand_passengers"] == 0
    assert buckets[0]["interval_start"] == 300
    assert buckets[0]["interval_end"] == 330
    assert {"A2", "D2", "D3", "K2", "L2", "K3", "L3"} <= set(buckets[0]["source_cells"])
    assert "RECONCILED" in buckets[0]["normalization_method"]


@pytest.mark.parametrize(("cell", "value"), [("D3", 1), ("L4", 99), ("A4", "05:45"), ("I3", None)])
def test_ambiguous_demand_withholds_metrics_instead_of_selecting_convention(tmp_path, cell, value):
    data = external_import.import_external_workbook(
        changed(tmp_path, "Du lieu bieu do", cell, value)
    )
    assert data["demand_status"] == "UNRESOLVED_DEMAND_AGGREGATION"
    assert data["demand"] == []
    assert any(i["code"] == "UNRESOLVED_DEMAND_AGGREGATION" for i in data["issues"])


def test_extracts_both_vehicle_tables_and_summary_separately():
    data = external_import.import_external_workbook(FIXTURE)
    assert len(data["vehicle_trips"]) == 102
    assert {t["direction"] for t in data["vehicle_trips"]} == {"outbound", "inbound"}
    assert data["vehicle_trips"][0]["arrival"] == 305
    assert data["vehicle_trips"][0]["recorded_layover"] == 10
    assert len(data["vehicle_summaries"]) == 2
    assert data["vehicle_summaries"][0]["reported_operating_duration"] == 255


def test_missing_inbound_vehicle_row_is_explicit_and_does_not_hide_timetable(tmp_path):
    data = external_import.import_external_workbook(
        changed(tmp_path, "Lich cho tai xe", "M2", None)
    )
    assert any(i["code"] == "INCOMPLETE_VEHICLE_ROW" for i in data["issues"])
    assert len(data["timetables"]["EXTERNAL_CANDIDATE"]["inbound"]) == 51


def test_excluded_chart_observations_are_preserved_with_source_cells(tmp_path):
    data = external_import.import_external_workbook(changed(tmp_path, "Du lieu bieu do", "M2", 12))
    assert data["demand_exclusions"][0] == {
        "direction": "outbound",
        "observation_time": 300,
        "demand_passengers": 12,
        "source_sheet": "Du lieu bieu do",
        "source_cells": ["K2", "M2"],
        "normalization_method": "SOURCE_OUTSIDE_FRAME_EXCLUDED_FROM_HALF_HOUR_SERIES",
    }
    assert (
        sum(b["demand_passengers"] for b in data["demand"] if b["direction"] == "outbound") == 500
    )


def test_wrong_demand_header_is_unresolved_instead_of_silently_reinterpreted(tmp_path):
    data = external_import.import_external_workbook(
        changed(tmp_path, "Du lieu bieu do", "D1", "Revenue")
    )
    assert data["demand"] == []


def test_normalized_hourly_and_directional_controls_keep_source_references():
    data = external_import.import_external_workbook(FIXTURE)
    control = data["timetable_controls"]["BASELINE"]["outbound"]
    assert control["directional_total"] == {"source_cell": "U6", "count": 51}
    assert control["hourly_totals"][1] == {"hour": 5, "source_cell": "C6", "count": 4}


def test_changed_vehicle_headers_do_not_silently_reassign_columns(tmp_path):
    data = external_import.import_external_workbook(
        changed(tmp_path, "Lich cho tai xe", "D1", "arrival_time")
    )
    assert any(i["code"] == "UNSUPPORTED_VEHICLE_HEADERS" for i in data["issues"])
    assert data["vehicle_trips"] == []


def test_truncated_demand_cannot_be_reported_as_full_service_day(tmp_path):
    workbook = openpyxl.load_workbook(FIXTURE)
    sheet = workbook["Du lieu bieu do"]
    for row in range(26, sheet.max_row + 1):
        for col in (1, 2, 3, 4, 6, 7, 8, 9, 11, 12, 13, 15, 16, 17):
            sheet.cell(row, col).value = None
    path = tmp_path / "partial.xlsx"
    workbook.save(path)
    workbook.close()
    data = external_import.import_external_workbook(path)
    assert data["demand_status"] == "UNRESOLVED_DEMAND_AGGREGATION"
    assert data["demand"] == []


def test_outside_frame_observation_without_timestamp_is_unresolved(tmp_path):
    data = external_import.import_external_workbook(changed(tmp_path, "Du lieu bieu do", "M54", 75))
    assert data["demand_status"] == "UNRESOLVED_DEMAND_AGGREGATION"
